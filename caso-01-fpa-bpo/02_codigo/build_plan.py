"""Planificador de presupuesto y forecast (oct-25 a dic-26).

Genera, desde 01_datos/dashboard_data.json:
  - plan_map.json      : base, supuestos por defecto y direcciones de celdas (lo lee el tablero)
  - plan_template.xlsx : Excel con fórmulas vivas (el tablero lo descarga con los supuestos que cargues)
  - ../03_tableros/AAAA-MM-DD_Planificador-2025-2026_vN.xlsx : el mismo Excel con los supuestos por defecto

Lógica (igual en el tablero y en el Excel):
  base = run-rate jul–sep 2025 a precios de sep-25 (ingresos desestacionalizados)
  cada línea de costo/ingreso se mueve con un índice según su categoría:
    rev / iibb : índice de precios (o de TC si la cartera factura en US$) × volumen × recupero/recupero base × fee/fee base × estacionalidad
    pos_sal    : índice salarial × posiciones / posiciones base          (operadores, supervisión y apoyo)
    pos_usd    : índice de TC × posiciones / posiciones base             (posiciones Paraguay, licencias)
    pos_inf    : índice de precios × posiciones / posiciones base        (telefonía)
    vol_usd    : índice de TC × volumen                                  (mailing, WhatsApp)
    vol_inf    : índice de precios × volumen                             (bocas de cobranza, informes)
    sms        : índice de precio SMS × volumen
    estructura : sal (salarial), usd (TC), inf (precios), fix (fijo), ic (salarial)
    difc       : saldo US$ intercompany × variación del TC del mes
    tax        : tasa × resultado antes de impuestos (si es positivo)
"""
import os, sys, json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

D = os.path.dirname(os.path.abspath(__file__))
DATA = json.load(open(os.path.join(D, "..", "01_datos", "dashboard_data.json"), encoding="utf-8"))
ver = sys.argv[1] if len(sys.argv) > 1 else "v1"
today = datetime.date.today().isoformat()

MESES = DATA["meses"]; CORTE = MESES.index(DATA["meta"]["corte"]); Y25 = MESES.index("2025-01"); DIC = MESES.index("2025-12")
IT = DATA["items"]; NI = len(IT); CAR = DATA["carteras"]; NC = len(CAR)
ITEM = {x["item"]: i for i, x in enumerate(IT)}
SEAS = [0.96, 0.97, 1.03, 1.01, 1.00, 0.99, 1.00, 1.01, 1.02, 1.01, 1.00, 0.97]
MES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]

# ---------- cubo e índices (igual que el tablero) ----------
cube = {}
R = DATA["res"]
for e, m, c, i, v in zip(R["e"], R["m"], R["c"], R["i"], R["v"]):
    a = cube.setdefault((e, m), [[0.0] * NI for _ in range(NC + 1)])
    a[c + 1][i] += v
KP = {(r[0], r[1]): dict(tc=r[2], ops=r[3], py=r[4], pos=r[5], horas=r[6], cop=r[7], csms=r[8], cmin=r[9], I=r[10], P=r[11]) for r in DATA["kpi"]}
DR = {(r[0], r[1], r[2]): r for r in DATA["drv"]}
GRPN = {}
for n in ["Ingreso por recupero", "Comisiones por venta"]: GRPN[n] = "rev"
GRPN["Ingresos Brutos"] = "iibb"
for n in ["Operadores", "Supervisión", "Calidad", "Back Office", "Ejecutivos de cuenta", "Gerencia de Operaciones", "Indemnizaciones", "Remuneraciones staff"]: GRPN[n] = "rem"
for n in ["Posiciones Paraguay", "Licencias discador/CRM", "Mailing", "WhatsApp", "Infraestructura TI (USD)"]: GRPN[n] = "usd"
GRPN["Amortizaciones"] = "fix"; GRPN["Diferencia de cambio"] = "difc"; GRPN["Recupero gastos intercompany"] = "ic"; GRPN["Impuesto a las ganancias"] = "tax"
GRP = [GRPN.get(x["item"], "inf") for x in IT]
isUSD = lambda ci: ci >= 0 and CAR[ci]["moneda"] == "USD"

# run-rate jul–sep a precios de sep (ingresos desestacionalizados)
r0 = KP[(0, CORTE)]
RR = [[0.0] * NI for _ in range(NC + 1)]
for m in [CORTE - 2, CORTE - 1, CORTE]:
    k = KP[(0, m)]; a = cube[(0, m)]; seas = SEAS[int(MESES[m][5:]) - 1]
    for cc in range(NC + 1):
        ci = cc - 1
        for i in range(NI):
            v = a[cc][i]
            if not v: continue
            g = GRP[i]
            if g in ("rev", "iibb"): df = ((r0["tc"] / k["tc"]) if isUSD(ci) else (r0["I"] / k["I"])) / seas
            elif g in ("rem", "ic"): df = r0["P"] / k["P"]
            elif g == "usd": df = r0["tc"] / k["tc"]
            elif g in ("fix", "difc", "tax"): df = 1
            else: df = r0["I"] / k["I"]
            RR[cc][i] += v * df / 3

# categoría por línea
CATC = {"Ingreso por recupero": "rev", "Comisiones por venta": "rev", "Ingresos Brutos": "iibb", "Indemnizaciones": "zero",
        "Operadores": "pos_sal", "Supervisión": "pos_sal", "Calidad": "pos_sal", "Back Office": "pos_sal", "Ejecutivos de cuenta": "pos_sal", "Gerencia de Operaciones": "pos_sal",
        "Posiciones Paraguay": "pos_usd", "Licencias discador/CRM": "pos_usd", "Mailing": "vol_usd", "WhatsApp": "vol_usd", "SMS": "sms",
        "Telefonía celular": "pos_inf", "Telefonía LDN": "pos_inf", "Telefonía local y 0800": "pos_inf"}
def cat_of(cc, i):
    n = IT[i]["item"]; g = GRP[i]
    if cc == 0:  # estructura
        return {"rem": "sal", "usd": "usd", "fix": "fix", "ic": "ic", "difc": "difc", "tax": "tax", "rev": "inf", "iibb": "inf"}.get(g, "inf")
    if n in CATC: return CATC[n]
    return {"rem": "pos_sal", "usd": "vol_usd", "fix": "fix", "difc": "difc", "tax": "tax"}.get(g, "vol_inf")

CAT_TXT = {"rev": "Ingreso: precios (o TC) × volumen × recupero × fee × estacionalidad", "iibb": "Igual que el ingreso (es un % de lo facturado)",
           "pos_sal": "Índice salarial × posiciones", "pos_usd": "Índice de TC × posiciones", "pos_inf": "Índice de precios × posiciones",
           "vol_usd": "Índice de TC × volumen", "vol_inf": "Índice de precios × volumen", "sms": "Índice de precio SMS × volumen",
           "zero": "No se presupuesta (se carga aparte si hay desvinculaciones)", "sal": "Índice salarial", "usd": "Índice de TC", "inf": "Índice de precios",
           "fix": "Fijo", "ic": "Índice salarial", "difc": "Saldo US$ × variación del TC", "tax": "Tasa × resultado antes de impuestos"}

# ---------- horizonte, supuestos por defecto ----------
T = 15  # oct-25 .. dic-26
mlab = [MES[(9 + t) % 12] + "-" + ("25" if t < 3 else "26") for t in range(T)]
moidx = [(9 + t) % 12 for t in range(T)]
tcSep = r0["tc"]; tcDicF = KP[(2, DIC)]["tc"]
infl = [0.023] * 3 + [0.019] * 12
tc = [round(tcSep + (tcDicF - tcSep) * (t + 1) / 3) for t in range(3)] + [round(tcDicF + (1750 - tcDicF) * (k + 1) / 12) for k in range(12)]
par = [0.0] * T
for t, lab in enumerate(mlab):
    if lab in ("oct-25",): par[t] = 0.06
    if lab in ("ene-26", "abr-26", "jul-26", "oct-26"): par[t] = 0.055
sms = list(infl)
seas = [SEAS[mo] for mo in moidx]
icr = [r for r in DATA["ic"] if r[0] == 0 and CORTE - 2 <= r[1] <= CORTE]
saldo = sum(r[5] for r in icr) / len(icr)
tasa = 0.35

# drivers base por cartera (sep-25 real)
cars = []
for ci, c in enumerate(CAR):
    r = DR.get((0, CORTE, ci))
    if not r or not r[3]: continue
    cars.append(dict(ci=ci, code=c["code"], nombre=c["nombre"], mon=c["moneda"], recB=r[5], feeB=r[6], posB=r[8] + r[9]))
carIdx = {c["ci"]: c for c in cars}

# filas de cálculo (ámbito × ítem con base distinta de cero), impuesto al final
rows = []
for cc in list(range(1, NC + 1)) + [0]:
    ci = cc - 1
    if ci >= 0 and ci not in carIdx: continue
    for i in range(NI):
        cat = cat_of(cc, i)
        if cat == "tax": continue
        if abs(RR[cc][i]) < 1 and not (cc == 0 and cat == "difc"): continue
        rows.append(dict(cc=cc, ci=ci, i=i, amb=(CAR[ci]["code"] if ci >= 0 else "Estructura"), mon=(CAR[ci]["moneda"] if ci >= 0 else "ARS"),
                         n1=IT[i]["n1"], n2=IT[i]["n2"], item=IT[i]["item"], cat=cat, base=round(RR[cc][i], 2)))
rows.append(dict(cc=0, ci=-1, i=ITEM["Impuesto a las ganancias"], amb="Estructura", mon="ARS", n1=IT[ITEM["Impuesto a las ganancias"]]["n1"],
                 n2=IT[ITEM["Impuesto a las ganancias"]]["n2"], item="Impuesto a las ganancias", cat="tax", base=0))

# real ene–sep 2025 y presupuesto 2025 por N2 (para el cierre)
N2L = []
for x in IT:
    if x["n2"] not in N2L: N2L.append(x["n2"])
N1L = []
for x in IT:
    if x["n1"] not in N1L: N1L.append(x["n1"])
real25 = {n: 0.0 for n in N2L}; ppto25 = {n: 0.0 for n in N2L}
for (e, m), a in cube.items():
    for cc in range(NC + 1):
        for i in range(NI):
            if e == 0 and Y25 <= m <= CORTE: real25[IT[i]["n2"]] += a[cc][i]
            if e == 1 and Y25 <= m <= DIC: ppto25[IT[i]["n2"]] += a[cc][i]

# ---------- Excel ----------
wb = Workbook(); wb.calculation.fullCalcOnLoad = True
AR = "Arial"
F = lambda **k: Font(name=AR, size=k.pop("size", 10), **k)
NAVY = "13294B"; HDR = PatternFill("solid", fgColor=NAVY); INF = PatternFill("solid", fgColor="FFF2CC"); SUB = PatternFill("solid", fgColor="EEF2F7"); BLK = PatternFill("solid", fgColor="F5F7FA")
thin = Side(style="thin", color="D5DAE0"); BOT = Border(bottom=thin); TOP = Border(top=Side(style="thin", color="8B949E"))
F_IN = F(color="0000FF"); F_LINK = F(color="008000"); F_HDR = F(bold=True, color="FFFFFF"); F_B = F(bold=True)
MILL = '#,##0.0,,;(#,##0.0,,);"-"'; PCT = '0.0%;(0.0%);"-"'; NUM = '#,##0;(#,##0);"-"'; IDX = '0.0000'
def title(ws, t, sub):
    ws["A1"] = t; ws["A1"].font = F(size=14, bold=True, color=NAVY)
    ws["A2"] = sub; ws["A2"].font = F(color="5D6874")
def header(ws, row, vals, c0=1):
    for k, v in enumerate(vals):
        c = ws.cell(row=row, column=c0 + k, value=v); c.font = F_HDR; c.fill = HDR; c.alignment = Alignment(horizontal="center" if k else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30

# Leeme
ws0 = wb.active; ws0.title = "Leeme"
title(ws0, "Planificador · reproyección oct–dic 2025 y presupuesto 2026", f"Cobranzas del Sur S.A. · datos 100% ficticios · generado el {today}")
txt = [
    ("Para qué sirve", "Arma el cierre 2025 (real ene–sep + reproyección oct–dic) y el presupuesto 2026 mes a mes, desde la operación real de jul–sep 2025."),
    ("Qué se carga", "Celdas azules con fondo amarillo: supuestos mes a mes (hoja Supuestos) y lo que proyecta cada supervisor por cartera (hoja Drivers). Todo lo demás son fórmulas."),
    ("Supuestos", "Inflación mensual, tipo de cambio, aumentos salariales por tramos (paritarias), aumento del precio del SMS y estacionalidad."),
    ("Drivers por cartera", "Volumen de deuda asignada (índice real, sep-25 = 100), % de recupero, fee y posiciones. Para dar de baja una cartera, poné volumen y posiciones en 0."),
    ("Cálculo", "Una fila por cartera y línea del estado de resultados. La columna Categoría dice con qué índice se mueve cada línea (ver la lista abajo)."),
    ("EERR", "Estado de resultados mensual: real ene–sep, reproyección oct–dic, cierre 2025 contra presupuesto 2025, y presupuesto 2026 mes a mes."),
    ("Por cartera", "Ingresos y margen de contribución por cartera para el cierre 2025 y el presupuesto 2026."),
    ("Dotación", "Posiciones por cartera y remuneraciones operativas. Información sensible: solo RRHH, Gerencia y Control de Gestión."),
    ("Base", "Run-rate de jul–sep 2025 llevado a precios de septiembre; los ingresos se desestacionalizan y la estacionalidad se vuelve a aplicar mes a mes."),
    ("Montos", "En pesos; las tablas de resultados se muestran en millones (formato de celda). Índices: sep-25 = 1."),
    ("Al abrir", "Si Excel abre en Vista protegida y ves celdas vacías, tocá «Habilitar edición»: el archivo recalcula todo al abrir."),
]
for k, (a, b) in enumerate(txt):
    ws0.cell(row=4 + k, column=1, value=a).font = F_B; c = ws0.cell(row=4 + k, column=2, value=b); c.font = F(); c.alignment = Alignment(wrap_text=True, vertical="top")
r = 4 + len(txt) + 1
ws0.cell(row=r, column=1, value="Categorías de cálculo").font = F(bold=True, color=NAVY); r += 1
for k, v in CAT_TXT.items():
    ws0.cell(row=r, column=1, value=k).font = F(); ws0.cell(row=r, column=2, value=v).font = F(); r += 1
r += 1
ws0.cell(row=r, column=1, value="Colores").font = F(bold=True, color=NAVY); r += 1
c = ws0.cell(row=r, column=1, value="1.234"); c.font = F_IN; c.fill = INF; ws0.cell(row=r, column=2, value="Dato que se carga (supuesto o proyección)").font = F(); r += 1
ws0.cell(row=r, column=1, value="1.234").font = F(); ws0.cell(row=r, column=2, value="Fórmula").font = F(); r += 1
ws0.cell(row=r, column=1, value="1.234").font = F_LINK; ws0.cell(row=r, column=2, value="Vínculo a otra hoja").font = F()
ws0.column_dimensions["A"].width = 24; ws0.column_dimensions["B"].width = 110

# Supuestos
ws = wb.create_sheet("Supuestos")
title(ws, "Supuestos mes a mes", "Celdas azules: se cargan. Índices con base sep-25 = 1.")
MC = [L(4 + t) for t in range(T)]  # D..R
header(ws, 4, ["Supuesto", "Unidad", "Base sep-25"] + mlab + ["Comentario"])
S_ROW = dict(infl=5, I=6, tc=7, tR=8, par=9, S=10, sms=11, Sm=12, seas=13, saldo=16, tasa=17)
labels = [(5, "Inflación mensual", "%", None, infl, PCT, "Variación mensual de precios. Ajusta la deuda en pesos (y el ingreso) y los costos que siguen a la inflación."),
          (6, "Índice de precios", "índice", 1, None, IDX, "Acumula la inflación desde sep-25."),
          (7, "Tipo de cambio promedio", "$ / US$", tcSep, tc, NUM, "Mueve los ingresos de carteras en dólares y los costos en dólares (Paraguay, licencias, mailing)."),
          (8, "Índice de tipo de cambio", "índice", 1, None, IDX, "TC del mes / TC de sep-25."),
          (9, "Aumento salarial del mes (paritaria)", "%", None, par, PCT, "Cargá el % en el mes del tramo (trimestral, bimestral o mensual)."),
          (10, "Índice salarial", "índice", 1, None, IDX, "Acumula los tramos de paritaria."),
          (11, "Aumento del precio del SMS", "%", None, sms, PCT, "Por defecto, igual a la inflación."),
          (12, "Índice de precio SMS", "índice", 1, None, IDX, ""),
          (13, "Estacionalidad de ingresos", "factor", None, seas, '0.00', "Factor del mes calendario (1 = mes promedio).")]
for row, lab, un, base, vals, fmt, note in labels:
    ws.cell(row=row, column=1, value=lab).font = F_B if vals is None else F(); ws.cell(row=row, column=2, value=un).font = F(color="5D6874")
    if base is not None:
        c = ws.cell(row=row, column=3, value=base); c.number_format = fmt; c.font = F_IN if row == 7 else F(); c.fill = INF if row == 7 else PatternFill()
    for t in range(T):
        col = 4 + t; c = ws.cell(row=row, column=col)
        if vals is not None: c.value = vals[t]; c.font = F_IN; c.fill = INF
        else:
            prev = L(col - 1)
            if row == 6: c.value = f"={prev}6*(1+{L(col)}5)"
            elif row == 8: c.value = f"={L(col)}7/$C$7"
            elif row == 10: c.value = f"={prev}10*(1+{L(col)}9)"
            elif row == 12: c.value = f"={prev}12*(1+{L(col)}11)"
            c.font = F()
        c.number_format = fmt
    ws.cell(row=row, column=4 + T, value=note).font = F(color="5D6874", size=9)
ws.cell(row=15, column=1, value="Parámetros").font = F(bold=True, color=NAVY)
ws.cell(row=16, column=1, value="Saldo a cobrar intercompany").font = F(); ws.cell(row=16, column=2, value="US$").font = F(color="5D6874")
c = ws.cell(row=16, column=3, value=round(saldo)); c.font = F_IN; c.fill = INF; c.number_format = NUM
ws.cell(row=16, column=4, value="Promedio jul–sep 2025. La diferencia de cambio del mes = saldo × variación del TC.").font = F(color="5D6874", size=9)
ws.cell(row=17, column=1, value="Tasa de impuesto a las ganancias").font = F(); ws.cell(row=17, column=2, value="%").font = F(color="5D6874")
c = ws.cell(row=17, column=3, value=tasa); c.font = F_IN; c.fill = INF; c.number_format = PCT
ws.cell(row=7, column=3).comment = Comment("Tipo de cambio promedio real de septiembre 2025 (fuente: KPI de operación del caso).", "Planificador")
ws.column_dimensions["A"].width = 36; ws.column_dimensions["B"].width = 10; ws.column_dimensions["C"].width = 12
for t in range(T): ws.column_dimensions[MC[t]].width = 10.5
ws.column_dimensions[L(4 + T)].width = 70
ws.freeze_panes = "D5"

# Drivers
wd = wb.create_sheet("Drivers")
title(wd, "Drivers por cartera · lo que proyecta cada supervisor", "Volumen: índice real (sep-25 = 100). Recupero y fee en %. Posiciones: operadores en Argentina + posiciones en Paraguay.")
header(wd, 4, ["Cartera / driver", "Unidad", "Base sep-25"] + mlab)
DROWS = {}
r = 5
for c in cars:
    wd.cell(row=r, column=1, value=f"{c['code']} · {c['nombre']} ({c['mon']})").font = F(bold=True, color=NAVY)
    for col in range(1, 4 + T): wd.cell(row=r, column=col).fill = SUB
    spec = [("vol", "Volumen asignado", "índice", 100, '0.0'), ("rec", "% de recupero", "%", c["recB"], '0.00%'), ("fee", "Fee", "%", c["feeB"], '0.00%'), ("pos", "Posiciones", "n", c["posB"], '0')]
    DROWS[c["code"]] = {}
    for k, (key, lab, un, base, fmt) in enumerate(spec):
        rr = r + 1 + k; DROWS[c["code"]][key] = rr
        wd.cell(row=rr, column=1, value="   " + lab).font = F(); wd.cell(row=rr, column=2, value=un).font = F(color="5D6874")
        cb = wd.cell(row=rr, column=3, value=base); cb.number_format = fmt; cb.font = F()
        for t in range(T):
            cc_ = wd.cell(row=rr, column=4 + t, value=base); cc_.font = F_IN; cc_.fill = INF; cc_.number_format = fmt
    r += 6
wd.column_dimensions["A"].width = 44; wd.column_dimensions["B"].width = 9; wd.column_dimensions["C"].width = 12
for t in range(T): wd.column_dimensions[MC[t]].width = 10
wd.freeze_panes = "D5"

# Cálculo
wc = wb.create_sheet("Cálculo")
title(wc, "Cálculo por cartera y línea", "Cada celda = base mensual × índices de la categoría. No se carga nada acá.")
CM = [L(8 + t) for t in range(T)]  # H..V
header(wc, 4, ["Ámbito", "Moneda", "N1", "N2", "Ítem", "Categoría", "Base mensual sep-25 ($)"] + mlab)
first = 5; last = first + len(rows) - 1
for k, rw in enumerate(rows):
    rr = first + k
    for col, v in enumerate([rw["amb"], rw["mon"], rw["n1"], rw["n2"], rw["item"], rw["cat"]], start=1): wc.cell(row=rr, column=col, value=v).font = F()
    cb = wc.cell(row=rr, column=7, value=rw["base"]); cb.number_format = NUM; cb.font = F()
    dr_ = DROWS.get(rw["amb"])
    for t in range(T):
        sc = MC[t]; S_ = lambda key: f"Supuestos!{sc}${S_ROW[key]}"
        if dr_:
            V = f"(Drivers!{sc}${dr_['vol']}/100)"; P_ = f"(Drivers!{sc}${dr_['pos']}/Drivers!$C${dr_['pos']})"
            REC = f"(Drivers!{sc}${dr_['rec']}/Drivers!$C${dr_['rec']})"; FEE = f"(Drivers!{sc}${dr_['fee']}/Drivers!$C${dr_['fee']})"
        cat = rw["cat"]; G = f"$G{rr}"
        if cat in ("rev", "iibb"): f = f"={G}*{S_('tR' if rw['mon'] == 'USD' else 'I')}*{V}*{REC}*{FEE}*{S_('seas')}"
        elif cat == "pos_sal": f = f"={G}*{S_('S')}*{P_}"
        elif cat == "pos_usd": f = f"={G}*{S_('tR')}*{P_}"
        elif cat == "pos_inf": f = f"={G}*{S_('I')}*{P_}"
        elif cat == "vol_usd": f = f"={G}*{S_('tR')}*{V}"
        elif cat == "vol_inf": f = f"={G}*{S_('I')}*{V}"
        elif cat == "sms": f = f"={G}*{S_('Sm')}*{V}"
        elif cat == "zero": f = "=0"
        elif cat in ("sal", "ic"): f = f"={G}*{S_('S')}"
        elif cat == "usd": f = f"={G}*{S_('tR')}"
        elif cat == "inf": f = f"={G}*{S_('I')}"
        elif cat == "fix": f = f"={G}"
        elif cat == "difc": f = f"=Supuestos!$C$16*(Supuestos!{sc}$7-Supuestos!{L(3 + t)}$7)"
        elif cat == "tax": f = f"=-Supuestos!$C$17*MAX(0,SUM({CM[t]}{first}:{CM[t]}{last - 1}))"
        c = wc.cell(row=rr, column=8 + t, value=f); c.number_format = NUM; c.font = F()
    if rw["cat"] == "tax":
        for col in range(1, 8 + T): wc.cell(row=rr, column=col).border = TOP
for col, w in zip("ABCDEFG", [12, 8, 26, 28, 28, 10, 16]): wc.column_dimensions[col].width = w
for t in range(T): wc.column_dimensions[CM[t]].width = 13
wc.freeze_panes = "H5"; wc.auto_filter.ref = f"A4:{CM[-1]}{last}"
CR = lambda col: f"Cálculo!${col}${first}:${col}${last}"

# EERR
we = wb.create_sheet("EERR")
title(we, "Estado de resultados de gestión · cierre 2025 y presupuesto 2026", "En millones de pesos (formato de celda). Real ene–sep: datos de gestión; oct–dic: reproyección; 2026: presupuesto. Si ves celdas vacías, habilitá la edición.")
hdr = ["Línea", "Real ene–sep 2025"] + mlab[:3] + ["Cierre 2025", "Presupuesto 2025", "Cierre vs presup."] + mlab[3:] + ["Total 2026", "2026 vs cierre 2025"]
header(we, 4, hdr)
# columnas: A línea, B real, C..E oct-dic, F cierre, G ppto25, H dif, I..T 2026, U total, V var
monthColE = ["C", "D", "E"] + [L(9 + k) for k in range(12)]
r = 5; n1rows = {}; subrows = {}
def eerr_line(rr, label, kind, n2=None, refs=None):
    we.cell(row=rr, column=1, value=label)
    for t in range(T):
        col = monthColE[t]
        if kind == "n2": f = f"=SUMIFS({CR(CM[t])},{CR('D')},$A{rr})"
        else: f = "=" + "+".join(f"{col}{x}" for x in refs)
        c = we.cell(row=rr, column=we[col + '1'].column, value=f); c.number_format = MILL
    if kind == "n2":
        we.cell(row=rr, column=2, value=round(real25[n2], 2)); we.cell(row=rr, column=7, value=round(ppto25[n2], 2))
    else:
        we.cell(row=rr, column=2, value="=" + "+".join(f"B{x}" for x in refs)); we.cell(row=rr, column=7, value="=" + "+".join(f"G{x}" for x in refs))
    we.cell(row=rr, column=6, value=f"=B{rr}+SUM(C{rr}:E{rr})"); we.cell(row=rr, column=8, value=f"=F{rr}-G{rr}")
    we.cell(row=rr, column=21, value=f"=SUM(I{rr}:T{rr})"); we.cell(row=rr, column=22, value=f'=IFERROR(U{rr}/F{rr}-1,"-")')
    for col in [2, 6, 7, 8, 21]: we.cell(row=rr, column=col).number_format = MILL
    we.cell(row=rr, column=22).number_format = PCT
    font = F() if kind == "n2" else F_B
    for col in range(1, 23):
        c = we.cell(row=rr, column=col); c.font = font
        if kind == "n1": c.fill = BLK
        if kind == "sub": c.fill = SUB; c.border = TOP
def pct_line(rr, label, num, den):
    we.cell(row=rr, column=1, value=label).font = F(italic=True, color="5D6874")
    for col in [2] + list(range(3, 6)) + [6, 7] + list(range(9, 22)):
        cl = L(col); c = we.cell(row=rr, column=col, value=f"=IFERROR({cl}{num}/{cl}{den},0)"); c.number_format = PCT; c.font = F(italic=True, color="5D6874")
SUBDEF = {2: ("Margen bruto", [0, 1, 2]), 3: ("Margen operativo", [0, 1, 2, 3]), 4: ("Resultado antes de impuestos", [0, 1, 2, 3, 4]), 5: ("Utilidad neta", [0, 1, 2, 3, 4, 5])}
for k, n1 in enumerate(N1L):
    kids = [n for n in N2L if any(x["n2"] == n and x["n1"] == n1 for x in IT)]
    hdr_r = r; r += 1; krows = []
    for n2 in kids: eerr_line(r, n2, "n2", n2=n2); krows.append(r); r += 1
    eerr_line(hdr_r, n1, "n1", refs=krows); n1rows[k] = hdr_r
    if k in SUBDEF:
        lab, ks = SUBDEF[k]; eerr_line(r, lab, "sub", refs=[n1rows[j] for j in ks]); subrows[lab] = r; r += 1
        pct_line(r, "   % sobre ingresos", subrows[lab], n1rows[0]); r += 1
we.column_dimensions["A"].width = 34
for col in range(2, 23): we.column_dimensions[L(col)].width = 11.5
we.freeze_panes = "B5"

# Por cartera
wp = wb.create_sheet("Por cartera")
title(wp, "Por cartera · margen de contribución", "Ingresos y margen (ingresos menos costos directos y asignados de la cartera). En millones de pesos.")
header(wp, 4, ["Código", "Cartera", "Moneda", "Ingresos oct–dic 25", "Margen oct–dic 25", "Ingresos 2026", "Margen 2026", "Margen % 2026", "Posiciones promedio 2026", "Ingreso mensual por posición 2026"])
r = 5
rngQ4 = f"Cálculo!$H${first}:$J${last}"; rng26 = f"Cálculo!$K${first}:$V${last}"; A_ = f"Cálculo!$A${first}:$A${last}"; F_ = f"Cálculo!$F${first}:$F${last}"
for c in cars:
    wp.cell(row=r, column=1, value=c["code"]); wp.cell(row=r, column=2, value=c["nombre"]); wp.cell(row=r, column=3, value=c["mon"])
    wp.cell(row=r, column=4, value=f'=SUMPRODUCT(({A_}=$A{r})*({F_}="rev")*{rngQ4})')
    wp.cell(row=r, column=5, value=f"=SUMPRODUCT(({A_}=$A{r})*{rngQ4})")
    wp.cell(row=r, column=6, value=f'=SUMPRODUCT(({A_}=$A{r})*({F_}="rev")*{rng26})')
    wp.cell(row=r, column=7, value=f"=SUMPRODUCT(({A_}=$A{r})*{rng26})")
    wp.cell(row=r, column=8, value=f"=IFERROR(G{r}/F{r},0)")
    pr = DROWS[c["code"]]["pos"]
    wp.cell(row=r, column=9, value=f"=AVERAGE(Drivers!$G${pr}:$R${pr})")
    wp.cell(row=r, column=10, value=f"=IFERROR(F{r}/12/I{r},0)")
    for col in range(1, 11): wp.cell(row=r, column=col).font = F()
    for col in [4, 5, 6, 7, 10]: wp.cell(row=r, column=col).number_format = MILL
    wp.cell(row=r, column=8).number_format = PCT; wp.cell(row=r, column=9).number_format = '0.0'
    r += 1
wp.cell(row=r, column=2, value="Total carteras")
for col in [4, 5, 6, 7]: wp.cell(row=r, column=col, value=f"=SUM({L(col)}5:{L(col)}{r - 1})").number_format = MILL
wp.cell(row=r, column=8, value=f"=IFERROR(G{r}/F{r},0)").number_format = PCT
wp.cell(row=r, column=9, value=f"=SUM(I5:I{r - 1})").number_format = '0.0'
for col in range(1, 11): c = wp.cell(row=r, column=col); c.font = F_B; c.fill = SUB; c.border = TOP
for col, w in zip("ABCDEFGHIJ", [9, 32, 9, 16, 16, 14, 14, 12, 14, 16]): wp.column_dimensions[col].width = w
wp.freeze_panes = "C5"

# Dotación
wt = wb.create_sheet("Dotación")
title(wt, "Dotación y remuneraciones operativas", "Información sensible: solo RRHH, Gerencia y Control de Gestión.")
wt["A2"].font = F(bold=True, color="A82A2A")
header(wt, 4, ["Cartera", "", "Sep-25"] + mlab)
r = 5
for c in cars:
    pr = DROWS[c["code"]]["pos"]
    wt.cell(row=r, column=1, value=f"{c['code']} · {c['nombre']}").font = F()
    wt.cell(row=r, column=3, value=f"=Drivers!C{pr}").font = F_LINK
    for t in range(T):
        cc_ = wt.cell(row=r, column=4 + t, value=f"=Drivers!{MC[t]}{pr}"); cc_.font = F_LINK; cc_.number_format = '0'
    r += 1
tot = r; wt.cell(row=r, column=1, value="Total posiciones").font = F_B
for col in range(3, 4 + T):
    c = wt.cell(row=r, column=col, value=f"=SUM({L(col)}5:{L(col)}{r - 1})"); c.font = F_B; c.fill = SUB; c.border = TOP; c.number_format = '0'
r += 2
wt.cell(row=r, column=1, value="Remuneraciones operativas ($)").font = F_B
for t in range(T):
    c = wt.cell(row=r, column=4 + t, value=f'=-SUMIFS({CR(CM[t])},{CR("F")},"pos_sal")'); c.number_format = MILL; c.font = F()
rem = r; r += 1
wt.cell(row=r, column=1, value="Costo mensual por posición ($)").font = F()
for t in range(T):
    c = wt.cell(row=r, column=4 + t, value=f"=IFERROR({MC[t]}{rem}/{MC[t]}{tot},0)"); c.number_format = NUM; c.font = F()
r += 1
wt.cell(row=r, column=1, value="Índice salarial (sep-25 = 1)").font = F()
for t in range(T):
    c = wt.cell(row=r, column=4 + t, value=f"=Supuestos!{MC[t]}10"); c.number_format = IDX; c.font = F_LINK
wt.cell(row=r + 2, column=1, value="Remuneraciones operativas: operadores, supervisión, calidad, back office, ejecutivos de cuenta y gerencia de operaciones (categoría pos_sal), en millones.").font = F(color="5D6874", size=9)
wt.column_dimensions["A"].width = 40; wt.column_dimensions["B"].width = 2; wt.column_dimensions["C"].width = 9
for t in range(T): wt.column_dimensions[MC[t]].width = 10
wt.freeze_panes = "D5"

for w in wb.worksheets: w.sheet_view.showGridLines = False
wb.active = wb.worksheets.index(we)
tpl = os.path.join(D, "plan_template.xlsx"); wb.save(tpl)
print("template", os.path.getsize(tpl) // 1024, "KB;", len(rows), "filas de cálculo")

# mapa para el tablero
sheets = {w.title: f"xl/worksheets/sheet{k + 1}.xml" for k, w in enumerate(wb.worksheets)}
M = dict(months=mlab, moidx=moidx, tcSep=tcSep, saldo=round(saldo), tasa=tasa,
         defaults=dict(infl=infl, tc=tc, par=par, sms=sms), seas=seas,
         car=[dict(ci=c["ci"], code=c["code"], nombre=c["nombre"], mon=c["mon"], recB=c["recB"], feeB=c["feeB"], posB=c["posB"], rows=DROWS[c["code"]]) for c in cars],
         rows=[dict(a=rw["amb"], ci=rw["ci"], i=rw["i"], cat=rw["cat"], base=rw["base"], mon=rw["mon"]) for rw in rows],
         cells=dict(sheets=sheets, monthCols=MC, sup=dict(infl=5, tc=7, par=9, sms=11), saldo="C16", tasa="C17"))
json.dump(M, open(os.path.join(D, "plan_map.json"), "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
print("map ok", len(json.dumps(M)) // 1024, "KB")

# entregable con los supuestos por defecto (recalculado si está LibreOffice, para que abra con valores)
import shutil, subprocess
dst = os.path.join(D, "..", "03_tableros", f"{today}_Planificador-2025-2026_{ver}.xlsx"); shutil.copy(tpl, dst)
rc = "/mnt/skills/public/xlsx/scripts/recalc.py"
if os.path.exists(rc):
    r = subprocess.run([sys.executable, rc, dst, "120"], capture_output=True, text=True); print(r.stdout.strip()[-200:])
print("excel ->", os.path.relpath(dst, D))
