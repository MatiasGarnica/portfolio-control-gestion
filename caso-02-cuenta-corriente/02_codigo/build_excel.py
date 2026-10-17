"""Base en Excel para auditar el tablero: datos crudos + días calle con fórmulas + la hoja de la semana.
Sale en ../01_datos/Base-cuenta-corriente.xlsx. Después: python recalc.py (LibreOffice) para dejar los valores calculados."""
import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from hoja import armar, cargar, PARAMS, FOCOS, ESTADOS

D = os.path.dirname(os.path.abspath(__file__))
DAT = os.path.join(D, "..", "01_datos")
rd = lambda f, **k: pd.read_csv(os.path.join(DAT, f), **k)
SUC, SUP, REP, CLI = rd("dim_sucursal.csv"), rd("dim_supervisor.csv"), rd("dim_reparto.csv"), rd("dim_cliente.csv")
V, P = rd("fact_venta.csv", parse_dates=["fecha"]), rd("fact_pago.csv", parse_dates=["fecha"])
M, VI, GE = rd("fact_movimiento.csv", parse_dates=["fecha"]), rd("fact_visita.csv", parse_dates=["fecha"]), rd("fact_gestion.csv", parse_dates=["fecha"])
data = cargar()

com = set(REP[REP.tipo_reparto == "Comercial"].id_reparto)
CLI["en_alcance"] = np.where((CLI.cuenta_interna == 0) & CLI.id_reparto.isin(com), "Sí", "No")
alc = dict(zip(CLI.id_cliente, CLI.en_alcance))
for df in (V, P, M):
    df["mes"] = df.fecha.dt.strftime("%Y-%m")
    df["en_alcance"] = df.id_cliente.map(alc)

FONT = "Arial"
H_FILL = PatternFill("solid", fgColor="13294B")
H_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
B_FONT = Font(name=FONT, size=10)
IN_FONT = Font(name=FONT, size=10, color="0000FF")        # valores cargados (no fórmula)
LINK_FONT = Font(name=FONT, size=10, color="008000")      # fórmulas que leen otra hoja
THIN = Border(bottom=Side(style="thin", color="D0D5DB"))
ARS = '"$" #,##0;-"$" #,##0;"-"'

wb = Workbook()
ws = wb.active
ws.title = "Léeme"

def tabla(nombre, df, formatos=None, anchos=None):
    sh = wb.create_sheet(nombre)
    sh.append(list(df.columns))
    for row in df.itertuples(index=False):
        sh.append([(x.to_pydatetime() if isinstance(x, pd.Timestamp) else (x.item() if hasattr(x, "item") else x)) for x in row])
    for c in sh[1]:
        c.font, c.fill, c.alignment = H_FONT, H_FILL, Alignment(vertical="center", wrap_text=True)
    sh.row_dimensions[1].height = 30
    sh.freeze_panes = "A2"
    sh.auto_filter.ref = sh.dimensions
    for j, col in enumerate(df.columns, 1):
        L = get_column_letter(j)
        sh.column_dimensions[L].width = (anchos or {}).get(col, max(11, min(34, len(col) + 3)))
        f = (formatos or {}).get(col)
        if f:
            for (cell,) in sh.iter_rows(min_row=2, min_col=j, max_col=j):
                cell.number_format = f
    return sh

# ---------------------------------------------------------------- datos crudos
cli_out = CLI.merge(REP[["id_reparto", "id_supervisor", "tipo_reparto"]], on="id_reparto").merge(SUP[["id_supervisor", "supervisor"]], on="id_supervisor", how="left")
cli_out = cli_out.merge(SUC[["id_sucursal", "sucursal"]], on="id_sucursal")
cli_out = cli_out[["id_cliente", "nombre", "tipo_cliente", "condicion", "medio_habitual", "frecuencia", "dia_visita", "id_reparto", "tipo_reparto",
                   "supervisor", "sucursal", "cuenta_interna", "en_alcance", "fecha_alta", "lat", "lng"]].sort_values("id_cliente")
cli_out["dia_visita"] = cli_out.dia_visita.map(dict(enumerate(["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado"])))
cli_out.columns = ["Cliente", "Nombre", "Tipo", "Condición", "Medio habitual", "Frecuencia", "Día de visita", "Reparto", "Tipo de reparto",
                   "Supervisor", "Sucursal", "Cuenta interna", "En alcance", "Alta", "Latitud", "Longitud"]
cli_out["Supervisor"] = cli_out.Supervisor.fillna("—")
tabla("Clientes", cli_out)

ven = V[["fecha", "id_cliente", "id_reparto", "importe", "unidades_sin_cargo", "mes", "en_alcance"]].sort_values(["fecha", "id_cliente"])
ven.columns = ["Fecha", "Cliente", "Reparto", "Importe", "Unidades sin cargo", "Mes", "En alcance"]
sv = tabla("Ventas", ven, {"Fecha": "dd/mm/yyyy", "Importe": ARS})
pag = P[["fecha", "id_cliente", "id_reparto", "importe", "medio", "mes", "en_alcance"]].sort_values(["fecha", "id_cliente"])
pag.columns = ["Fecha", "Cliente", "Reparto", "Importe", "Medio", "Mes", "En alcance"]
sp = tabla("Pagos", pag, {"Fecha": "dd/mm/yyyy", "Importe": ARS})
mov = M[["fecha", "id_cliente", "id_reparto", "tipo", "importe", "motivo", "mes", "en_alcance"]].sort_values(["fecha", "id_cliente"])
mov.columns = ["Fecha", "Cliente", "Reparto", "Tipo", "Importe", "Motivo", "Mes", "En alcance"]
sm = tabla("Movimientos", mov, {"Fecha": "dd/mm/yyyy", "Importe": ARS}, {"Motivo": 34, "Tipo": 18})
vis = VI.sort_values(["fecha", "id_cliente"]).copy()
vis["visito"] = np.where(vis.visito == 1, "Sí", "No")
vis.columns = ["Fecha", "Cliente", "Reparto", "Pasó el reparto"]
tabla("Visitas", vis, {"Fecha": "dd/mm/yyyy"})
ges = GE.sort_values("fecha").copy()
ges.columns = ["Fecha", "Cliente", "Área", "Gestión", "Resultado"]
tabla("Gestiones", ges, {"Fecha": "dd/mm/yyyy"}, {"Gestión": 28, "Resultado": 40})

# ---------------------------------------------------------------- días calle (fórmulas sobre las hojas de datos)
dc = wb.create_sheet("Días calle", 1)
nV, nP, nM = len(ven) + 1, len(pag) + 1, len(mov) + 1
cols = ["Reparto", "Supervisor", "Sucursal", "Mes", "Días de reparto", "Venta", "Deuda al cierre", "Días calle",
        "Cobrado total", "Cobrado en efectivo", "% efectivo", "Notas de crédito", "Incobrables", "NC + incobrables / venta"]
dc.append(cols)
sup_n = dict(zip(SUP.id_supervisor, SUP.supervisor)); suc_n = dict(zip(SUC.id_sucursal, SUC.sucursal))
rep_i = REP.set_index("id_reparto")
deuda = {(d["r"], m): d["m"][k][1] for d in data["dc"] for k, m in enumerate(data["meta"]["meses"])}
row = 2
for r in sorted(com):
    for m in data["meta"]["meses"]:
        dc.append([r, sup_n[rep_i.loc[r, "id_supervisor"]], suc_n[rep_i.loc[r, "id_sucursal"]], m, data["meta"]["dias_reparto"][m],
                   f'=SUMIFS(Ventas!$D$2:$D${nV},Ventas!$C$2:$C${nV},$A{row},Ventas!$F$2:$F${nV},$D{row},Ventas!$G$2:$G${nV},"Sí")',
                   deuda[(r, m)],
                   f'=IFERROR(G{row}/(F{row}/E{row}),"")',
                   f'=SUMIFS(Pagos!$D$2:$D${nP},Pagos!$C$2:$C${nP},$A{row},Pagos!$F$2:$F${nP},$D{row},Pagos!$G$2:$G${nP},"Sí")',
                   f'=SUMIFS(Pagos!$D$2:$D${nP},Pagos!$C$2:$C${nP},$A{row},Pagos!$F$2:$F${nP},$D{row},Pagos!$G$2:$G${nP},"Sí",Pagos!$E$2:$E${nP},"Efectivo")',
                   f'=IFERROR(J{row}/I{row},"")',
                   f'=SUMIFS(Movimientos!$E$2:$E${nM},Movimientos!$C$2:$C${nM},$A{row},Movimientos!$G$2:$G${nM},$D{row},Movimientos!$H$2:$H${nM},"Sí",Movimientos!$D$2:$D${nM},"Nota de crédito")',
                   f'=SUMIFS(Movimientos!$E$2:$E${nM},Movimientos!$C$2:$C${nM},$A{row},Movimientos!$G$2:$G${nM},$D{row},Movimientos!$H$2:$H${nM},"Sí",Movimientos!$D$2:$D${nM},"Incobrable")',
                   f'=IFERROR((L{row}+M{row})/F{row},"")'])
        row += 1
last = row - 1
for c in dc[1]:
    c.font, c.fill, c.alignment = H_FONT, H_FILL, Alignment(vertical="center", wrap_text=True)
dc.row_dimensions[1].height = 32
for rr in dc.iter_rows(min_row=2, max_row=last):
    for c in rr:
        c.font, c.border = B_FONT, THIN
    for j in (5, 7):
        rr[j - 1].font = IN_FONT
    for j in (6, 9, 10, 12, 13):
        rr[j - 1].font = LINK_FONT
    for j in (6, 7, 9, 10, 12, 13):
        rr[j - 1].number_format = ARS
    rr[7].number_format = '0.0;-0.0;"-"'
    rr[10].number_format = rr[13].number_format = '0.0%;-0.0%;"-"'
dc["G1"].comment = Comment("Valor calculado en build_data.py: saldo positivo de cada cuenta al último día del mes (o al corte), sumado por su reparto. "
                           "No se puede rearmar con una fórmula simple porque necesita el saldo de cada cliente a esa fecha.", "Matías Garnica")
dc["E1"].comment = Comment("Días de lunes a sábado del mes (octubre: hasta el corte del 09/10).", "Matías Garnica")
# total empresa por mes, al pie
tot0 = last + 3
dc.cell(tot0 - 1, 1, "Total empresa por mes").font = Font(name=FONT, bold=True, size=11)
dc.append([])
for k, m in enumerate(data["meta"]["meses"]):
    r_ = tot0 + k
    dc.cell(r_, 1, "Total"); dc.cell(r_, 4, m); dc.cell(r_, 5, data["meta"]["dias_reparto"][m]).font = IN_FONT
    for j, L in ((6, "F"), (7, "G"), (9, "I"), (10, "J"), (12, "L"), (13, "M")):
        c = dc.cell(r_, j, f"=SUMIFS({L}$2:{L}${last},$D$2:$D${last},$D{r_})"); c.number_format = ARS
    dc.cell(r_, 8, f'=IFERROR(G{r_}/(F{r_}/E{r_}),"")').number_format = '0.0;-0.0;"-"'
    dc.cell(r_, 11, f'=IFERROR(J{r_}/I{r_},"")').number_format = '0.0%;-0.0%;"-"'
    dc.cell(r_, 14, f'=IFERROR((L{r_}+M{r_})/F{r_},"")').number_format = '0.0%;-0.0%;"-"'
    for j in range(1, 15):
        c = dc.cell(r_, j); c.font = Font(name=FONT, bold=True, size=10, color="0000FF" if j == 5 else "000000"); c.border = THIN
dc.freeze_panes = "B2"
dc.auto_filter.ref = f"A1:N{last}"
for j, w in enumerate([10, 13, 10, 10, 10, 15, 15, 10, 15, 15, 10, 14, 13, 14], 1):
    dc.column_dimensions[get_column_letter(j)].width = w

# ---------------------------------------------------------------- hoja de la semana (último lunes, con el mismo motor)
W = len(data["meta"]["semanas"]) - 1
sem = data["meta"]["semanas"][W]
foco, est, nueva, cob = armar(data, W)
Cc = {c: i for i, c in enumerate(data["cols"])}
F = np.array(data["f"][W])
k = data["meta"]["umbrales"].index(PARAMS["dias"])
rep_sup = {r[0]: r[1] for r in data["rep"]}
FOCO_T = ["Deuda vencida", "Sin visita del reparto", "Ritmo de pago atrasado", "NC fuera de lo normal"]
EST_T = ["Resuelta", "En parte", "Pendiente", "Salió sola"]
rows = []
for i, c in enumerate(data["cli"]):
    if foco[i] < 0:
        continue
    r = F[i]
    rows.append([c[0], c[2], c[3], sup_n[rep_sup[c[3]]], FOCO_T[foco[i]], EST_T[est[i]], r[Cc[["v30", "v45", "v60", "v90"][k]]], r[Cc["saldo"]],
                 r[Cc["dup"]] if r[Cc["dup"]] >= 0 else None, r[Cc["imed"]] if r[Cc["imed"]] > 0 else None, r[Cc["duv"]] if r[Cc["duv"]] >= 0 else None,
                 r[Cc["nc30"]], r[Cc["pag"]], "Sí" if r[Cc["vis"]] else "No", int(r[Cc["lla"]])])
hj = pd.DataFrame(rows, columns=["Cliente", "Tipo", "Reparto", "Supervisor", "Foco", "Estado", f"Vencido a más de {PARAMS['dias']} días", "Saldo al domingo",
                                 "Días desde el último pago", "Intervalo habitual entre pagos", "Días sin visita", "NC último mes", "Pagó en la semana",
                                 "Pasó el reparto", "Llamados"]).sort_values(f"Vencido a más de {PARAMS['dias']} días", ascending=False)
sh = tabla(f"Hoja sem {sem['sem']}", hj, {f"Vencido a más de {PARAMS['dias']} días": ARS, "Saldo al domingo": ARS, "NC último mes": ARS, "Pagó en la semana": ARS})
wb.move_sheet(sh, -(len(wb.sheetnames) - 3))

# ---------------------------------------------------------------- léeme
txt = [
    ("Base de cuenta corriente · Reparto Federal S.A.", "t"),
    ("Datos 100% ficticios, generados con generate_case.py (semilla fija). Sirve para auditar cualquier número del tablero.", ""),
    ("", ""),
    ("Hojas", "h"),
    ("Días calle: venta, cobro, NC e incobrables por reparto y mes, con fórmulas SUMIFS sobre las hojas de datos. Días calle = deuda ÷ (venta ÷ días de reparto).", ""),
    (f"Hoja sem {sem['sem']}: la hoja de trabajo del lunes {sem['lunes'][8:10]}/{sem['lunes'][5:7]} con su foco y cómo va cada cuenta al corte (mismo criterio que hoja.py y el tablero).", ""),
    ("Clientes, Ventas, Pagos, Movimientos, Visitas, Gestiones: los datos crudos, uno por fila.", ""),
    ("", ""),
    ("Reglas de alcance", "h"),
    ("Las cuentas internas y las de repartos que no venden (técnicos y administración) quedan fuera: columna «En alcance» = No.", ""),
    ("La venta es solo con cargo: las unidades sin cargo de la promo 2×1 se informan aparte y no suman importe.", ""),
    ("Movimientos: notas de crédito, incobrables, reversas de pago (pago desconocido por el cliente) y el saldo inicial al 31/12/2024.", ""),
    ("", ""),
    ("Colores", "h"),
    ("Azul: valor cargado (no es fórmula). Verde: fórmula que lee otra hoja. Negro: fórmula de la misma hoja.", ""),
    (f"Parámetros acordados de la hoja: vencido a más de {PARAMS['dias']} días, piso $ {PARAMS['piso']:,}, sin visita más de {PARAMS['sin_visita']} días, "
     f"ritmo {PARAMS['factor']:g}× el intervalo habitual, NC del mes ≥ $ {PARAMS['piso_nc']:,} y > {PARAMS['pct_nc']:.0%} de la facturación mensual habitual.".replace(",", "."), ""),
]
for t, s in txt:
    ws.append([t])
    c = ws.cell(ws.max_row, 1)
    c.font = Font(name=FONT, bold=True, size=15, color="13294B") if s == "t" else Font(name=FONT, bold=True, size=11) if s == "h" else B_FONT
ws.column_dimensions["A"].width = 150
ws.sheet_view.showGridLines = False

wb.properties.creator = "Matías Garnica"
wb.properties.title = "Base de cuenta corriente · Reparto Federal S.A. (datos ficticios)"
path = os.path.join(DAT, "Base-cuenta-corriente.xlsx")
wb.save(path)
print("ok ->", os.path.relpath(path, D), os.path.getsize(path) // 1024, "KB ·", wb.sheetnames)
