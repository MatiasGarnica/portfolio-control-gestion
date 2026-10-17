"""Base en Excel del Caso 04: la base unificada (Excel limpio + formulario) con el arqueo, el resumen mensual y el fiado
calculados con fórmulas, y el registro de limpieza. Sale en ../01_datos/Base-pastas-don-aldo.xlsx.
Después: recalcular con LibreOffice (recalc.py) para dejar los valores."""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
import indicadores as I

c, v, p, cl = I.cargar()
P = pd.read_csv(os.path.join(I.BASE, "problemas_excel.csv"))
FONT = "Arial"
H_FILL, H_FONT = PatternFill("solid", fgColor="7C2D1B"), Font(name=FONT, bold=True, color="FFFFFF", size=10)
B, IN, LK, BOLD = Font(name=FONT, size=10), Font(name=FONT, size=10, color="0000FF"), Font(name=FONT, size=10, color="008000"), Font(name=FONT, size=10, bold=True)
THIN = Border(bottom=Side(style="thin", color="E3D9CC"))
ARS = '"$" #,##0;-"$" #,##0;"-"'
LM = "'Léeme'"                                            # hoja de supuestos (con tilde: va entre comillas)
wb = Workbook()
ws0 = wb.active; ws0.title = "Léeme"

def encabezado(ws, cols, anchos):
    ws.append(cols)
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = H_FONT, H_FILL, Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    for j, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = w

# ---------------------------------------------------------------- cierres con el arqueo en fórmulas
ws = wb.create_sheet("Cierres")
cols = ["Fecha", "Local", "Responsable", "Fuente", "Efectivo", "Débito", "Crédito", "Billetera virtual", "Transferencia", "Venta total", "Tickets",
        "Gasto con la caja", "Concepto", "Retiro del dueño", "Fondo", "Efectivo esperado", "Efectivo contado", "Diferencia", "Caja", "Incompleta"]
encabezado(ws, cols, [11, 9, 12, 11, 12, 11, 11, 12, 12, 13, 8, 11, 12, 12, 9, 13, 13, 11, 10, 10])
for i, r in enumerate(c.itertuples(), start=2):
    n = lambda x: None if pd.isna(x) else int(x)
    ws.append([r.fecha.to_pydatetime(), r.local, r.responsable if isinstance(r.responsable, str) else None, r.fuente, n(r.efectivo), int(r.debito), int(r.credito),
               int(r.billetera), int(r.transferencia), f"=SUM(E{i}:I{i})", n(r.tickets), int(r.gasto), r.concepto_gasto if isinstance(r.concepto_gasto, str) else None,
               int(r.retiro), n(r.fondo), f'=IF(O{i}="","",O{i}+E{i}-L{i}-N{i})', n(r.efectivo_contado), f'=IF(OR(Q{i}="",P{i}=""),"",Q{i}-P{i})',
               f'=IF(R{i}="","",IF(ABS(R{i})<={LM}!$B$12,"Cuadra",IF(R{i}<0,"Faltante","Sobrante")))', "Sí" if r.incompleta else "No"])
    for j, cell in enumerate(ws[i], 1):
        cell.font, cell.border = (IN if j in (5, 6, 7, 8, 9, 11, 12, 14, 15, 17) else LK if j == 19 else B), THIN
        if j in (5, 6, 7, 8, 9, 10, 12, 14, 15, 16, 17, 18):
            cell.number_format = ARS
    ws.cell(i, 1).number_format = "dd/mm/yyyy"
NC = len(c) + 1
ws.auto_filter.ref = f"A1:T{NC}"
ws["P1"].comment = Comment("Fondo + efectivo − gasto con la caja − retiro del dueño. Los cierres del Excel no tienen fondo ni efectivo contado.", "Matías Garnica")

# ---------------------------------------------------------------- resumen mensual por local (fórmulas sobre Cierres)
rs = wb.create_sheet("Resumen mensual", 1)
encabezado(rs, ["Mes", "Local", "Venta", "Efectivo", "% efectivo", "Tickets", "Ticket promedio", "Costo plataformas", "Cierres con arqueo", "Cierres que cuadran",
                "% que cuadra", "Faltó en la caja"], [10, 9, 14, 14, 10, 9, 12, 13, 11, 11, 10, 13])
meses = sorted(c.fecha.dt.strftime("%Y-%m").unique())
row = 2
rng = lambda col: f"Cierres!${col}$2:${col}${NC}"
for m in meses:
    y, mm = int(m[:4]), int(m[5:])
    ini, fin = f"DATE({y},{mm},1)", f"EOMONTH(DATE({y},{mm},1),0)"
    for l in ("Centro", "Jardín"):
        cond = f'{rng("A")},">="&{ini},{rng("A")},"<="&{fin},{rng("B")},$B{row}'
        rs.append([m, l, f"=SUMIFS({rng('J')},{cond})", f"=SUMIFS({rng('E')},{cond})", f'=IF(C{row}=0,"",D{row}/C{row})', f"=SUMIFS({rng('K')},{cond})",
                   f'=IF(F{row}=0,"",SUMIFS({rng("J")},{cond},{rng("K")},">0")/F{row})',
                   f"=ROUND(SUMIFS({rng('F')},{cond})*{LM}!$B$13+SUMIFS({rng('G')},{cond})*{LM}!$B$14+SUMIFS({rng('H')},{cond})*{LM}!$B$15,0)",
                   f'=COUNTIFS({cond},{rng("Q")},"<>")', f'=COUNTIFS({cond},{rng("S")},"Cuadra")', f'=IF(I{row}=0,"",J{row}/I{row})',
                   f'=-SUMIFS({rng("R")},{cond},{rng("R")},"<"&-{LM}!$B$12)'])
        for j, cell in enumerate(rs[row], 1):
            cell.font, cell.border = (LK if j >= 3 else B), THIN
            cell.number_format = ARS if j in (3, 4, 7, 8, 12) else "0.0%" if j in (5, 11) else "#,##0"
        row += 1
rs.freeze_panes = "C2"
rs.cell(row + 1, 1, "Fórmulas sobre la hoja Cierres. Hasta mayo (Excel) no hay tickets ni arqueo; hasta marzo la billetera está dentro de transferencia.").font = Font(name=FONT, size=9, italic=True)

# ---------------------------------------------------------------- pedidos, cobros, clientes y fiado
pe = wb.create_sheet("Pedidos")
encabezado(pe, ["Fecha", "Cliente", "Nombre", "Importe", "Condición", "Fuente"], [11, 8, 28, 12, 16, 11])
nom = dict(zip(cl.id_cliente, cl.nombre))
for r in v.itertuples():
    pe.append([r.fecha.to_pydatetime(), int(r.id_cliente), nom[r.id_cliente], int(r.importe), r.condicion, r.fuente])
co = wb.create_sheet("Cobros")
encabezado(co, ["Fecha", "Cliente", "Nombre", "Importe", "Medio", "Fuente"], [11, 8, 28, 12, 14, 11])
for r in p.itertuples():
    co.append([r.fecha.to_pydatetime(), int(r.id_cliente), nom[r.id_cliente], int(r.importe), r.medio, r.fuente])
for sh in (pe, co):
    for row_ in sh.iter_rows(min_row=2):
        row_[0].number_format, row_[3].number_format = "dd/mm/yyyy", ARS
        for cell in row_:
            cell.font = B
    sh.auto_filter.ref = sh.dimensions
NP, NB = len(v) + 1, len(p) + 1
fi = wb.create_sheet("Fiado", 2)
encabezado(fi, ["Cliente", "Nombre", "Condición", "Comprado a cuenta", "Cobrado", "Saldo", "Saldo del motor", "Control", "Comprado a cuenta (90 d)", "Cobrado (90 d)",
                "Pagó de lo comprado", "Estado"], [8, 28, 16, 15, 14, 14, 14, 10, 15, 14, 11, 22])
sal = I.saldo_mayoristas(v, p).set_index("id_cliente")
est = I.estado_clientes(v, cl).set_index("id_cliente")
d90 = f"{LM}!$B$11-90"
for k, r in enumerate(cl.itertuples(), start=2):
    s = sal.saldo.get(r.id_cliente)
    fi.append([int(r.id_cliente), r.nombre, r.condicion,
               f'=SUMIFS(Pedidos!$D$2:$D${NP},Pedidos!$B$2:$B${NP},A{k},Pedidos!$E$2:$E${NP},"Cuenta corriente")', f"=SUMIFS(Cobros!$D$2:$D${NB},Cobros!$B$2:$B${NB},A{k})",
               f'=IF(D{k}=0,"",D{k}-E{k})', None if s is None else int(s), f'=IF(G{k}="","",IF(F{k}=G{k},"OK","Revisar"))',
               f'=SUMIFS(Pedidos!$D$2:$D${NP},Pedidos!$B$2:$B${NP},A{k},Pedidos!$E$2:$E${NP},"Cuenta corriente",Pedidos!$A$2:$A${NP},">"&{d90})',
               f'=SUMIFS(Cobros!$D$2:$D${NB},Cobros!$B$2:$B${NB},A{k},Cobros!$A$2:$A${NB},">"&{d90})', f'=IF(I{k}=0,"",J{k}/I{k})',
               est.estado.get(r.id_cliente, "") + (" · paga menos" if s is not None and bool(sal.paga_menos.get(r.id_cliente)) else "")])
    for j, cell in enumerate(fi[k], 1):
        cell.font, cell.border = (LK if j in (4, 5, 9, 10) else IN if j == 7 else B), THIN
        cell.number_format = ARS if j in (4, 5, 6, 7, 9, 10) else "0.0%" if j == 11 else "General"
fi["G1"].comment = Comment("Saldo calculado por indicadores.py (igual que la vista v_saldo_mayoristas de Supabase). La columna Control lo compara con la fórmula.", "Matías Garnica")
fi.auto_filter.ref = f"A1:L{len(cl) + 1}"

pr = wb.create_sheet("Problemas del Excel")
encabezado(pr, ["Hoja", "Fila", "Problema", "Lo que decía", "Qué se hizo"], [10, 6, 30, 46, 44])
for r in P.itertuples():
    pr.append([r.hoja, None if pd.isna(r.fila) or r.fila == "" else r.fila, r.problema, r.detalle, r.que_se_hizo])
for row_ in pr.iter_rows(min_row=2):
    for cell in row_:
        cell.font = B
pr.auto_filter.ref = pr.dimensions

# ---------------------------------------------------------------- léeme y supuestos
txt = [("Base de caja y clientes · Pastas Don Aldo", "t"), ("Datos 100% ficticios (generate_pyme.py, semilla fija). Enero a mayo: Excel compartido limpiado con limpiar_excel.py. Junio en adelante: formulario.", ""), ("", ""),
       ("Hojas", "h"), ("Cierres: un cierre por local y día, con el arqueo en fórmulas (efectivo esperado, diferencia y si cuadra).", ""),
       ("Resumen mensual: venta, efectivo, tickets, costo de plataformas y caja por mes y local, con SUMIFS y COUNTIFS sobre Cierres.", ""),
       ("Fiado: lo comprado a cuenta menos lo cobrado, por restaurante; la columna Control compara la fórmula con el motor del tablero.", ""),
       ("Pedidos, Cobros y Problemas del Excel: los datos y el registro de cada arreglo de la limpieza.", ""), ("", ""), ("Supuestos (azul = valor cargado)", "h")]
for t, s in txt:
    ws0.append([t]); cc = ws0.cell(ws0.max_row, 1)
    cc.font = Font(name=FONT, bold=True, size=15, color="7C2D1B") if s == "t" else Font(name=FONT, bold=True, size=11) if s == "h" else B
sup = [("Corte de los datos", pd.Timestamp(I.CORTE).to_pydatetime(), "dd/mm/yyyy"), ("Tolerancia para que la caja cuadre ($)", I.TOL, ARS),
       ("Comisión débito", I.COMISION["debito"], "0.0%"), ("Comisión crédito", I.COMISION["credito"], "0.0%"), ("Comisión billetera virtual", I.COMISION["billetera"], "0.0%")]
for k, (t, val, f) in enumerate(sup, start=11):
    ws0.cell(k, 1, t).font = B
    cc = ws0.cell(k, 2, val); cc.font, cc.number_format = IN, f
ws0.append([]); ws0.append(["Colores: azul = valor cargado · verde = fórmula que lee otra hoja · negro = fórmula de la misma hoja."]); ws0.cell(ws0.max_row, 1).font = B
ws0.column_dimensions["A"].width = 110; ws0.column_dimensions["B"].width = 14
ws0.sheet_view.showGridLines = False
wb.properties.creator = "Matías Garnica"
wb.properties.title = "Base de caja y clientes · Pastas Don Aldo (datos ficticios)"
path = os.path.join(I.D, "..", "01_datos", "Base-pastas-don-aldo.xlsx")
wb.save(path)
print("ok ->", os.path.relpath(path, I.D), os.path.getsize(path) // 1024, "KB ·", wb.sheetnames)
