"""Base en Excel del Caso 03: las alertas del día y, para tres reglas, la condición comprobada con fórmulas sobre los datos crudos.
Sale en ../01_datos/Base-alertas.xlsx. Después: recalcular con LibreOffice (recalc.py) para dejar los valores."""
import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from alertas import cargar, candidatos, evaluar, PARAMS, REGLAS, CORTE, DIA, DAT

b = cargar()
C = candidatos(b)
A = evaluar(C)
cli = b["dim_cliente"].set_index("id_cliente")
FONT = "Arial"
H_FILL, H_FONT = PatternFill("solid", fgColor="13294B"), Font(name=FONT, bold=True, color="FFFFFF", size=10)
B_FONT, IN_FONT, LINK_FONT = Font(name=FONT, size=10), Font(name=FONT, size=10, color="0000FF"), Font(name=FONT, size=10, color="008000")
BOLD = Font(name=FONT, size=10, bold=True)
THIN = Border(bottom=Side(style="thin", color="D0D5DB"))
ARS, FECHA = '"$" #,##0;-"$" #,##0;"-"', "dd/mm/yyyy"
wb = Workbook()
ws0 = wb.active; ws0.title = "Léeme"

def hoja(nombre, cols, filas, formatos=None, anchos=None, fuentes=None):
    sh = wb.create_sheet(nombre)
    sh.append(cols)
    for f in filas:
        sh.append([x.to_pydatetime() if isinstance(x, pd.Timestamp) else (x.item() if hasattr(x, "item") else x) for x in f])
    for c in sh[1]:
        c.font, c.fill, c.alignment = H_FONT, H_FILL, Alignment(vertical="center", wrap_text=True)
    sh.row_dimensions[1].height = 30
    sh.freeze_panes = "A2"
    sh.auto_filter.ref = sh.dimensions
    for j, col in enumerate(cols, 1):
        sh.column_dimensions[get_column_letter(j)].width = (anchos or {}).get(col, max(11, min(40, len(col) + 3)))
        for (cell,) in sh.iter_rows(min_row=2, min_col=j, max_col=j):
            cell.font = (fuentes or {}).get(col, B_FONT)
            if (formatos or {}).get(col):
                cell.number_format = formatos[col]
    return sh

def df_hoja(nombre, df, formatos=None, anchos=None):
    return hoja(nombre, list(df.columns), df.itertuples(index=False), formatos, anchos)

# ---------------------------------------------------------------- datos crudos (los que usan las fórmulas)
ven = b["fact_venta"][["fecha", "id_cliente", "id_reparto", "importe", "unidades_sin_cargo"]]
ven.columns = ["Fecha", "Cliente", "Reparto", "Importe", "Unidades sin cargo"]
pag = b["fact_pago"][["fecha", "id_cliente", "id_reparto", "importe", "medio", "canal"]]
pag.columns = ["Fecha", "Cliente", "Reparto", "Importe", "Medio", "Canal"]
mov = b["fact_movimiento"][["fecha", "id_cliente", "id_reparto", "tipo", "importe", "motivo"]]
mov.columns = ["Fecha", "Cliente", "Reparto", "Tipo", "Importe", "Motivo"]
vis = b["fact_visita"].assign(visito=lambda x: np.where(x.visito == 1, "Sí", "No"))[["fecha", "id_cliente", "id_reparto", "visito"]]
vis.columns = ["Fecha", "Cliente", "Reparto", "Pasó el reparto"]
abo = b["fact_abono"][["fecha", "id_equipo", "id_cliente", "importe"]]
abo.columns = ["Fecha", "Equipo", "Cliente", "Importe"]
nV, nP, nM, nVi, nA = len(ven) + 1, len(pag) + 1, len(mov) + 1, len(vis) + 1, len(abo) + 1

# ---------------------------------------------------------------- resumen de alertas
cols = ["Regla", "Responsable", "Clave", "Cliente", "Reparto", "Se disparó", "En juego"]
rows = [[REGLAS[a.regla][0], REGLAS[a.regla][1], a.clave, a.id_cliente or None,
         a.id_reparto if isinstance(a.id_reparto, str) else cli.loc[a.id_cliente, "id_reparto"], a.fecha, a.impacto] for a in A.sort_values(["regla", "impacto"], ascending=[True, False]).itertuples()]
sh = hoja("Alertas", cols, rows, {"Se disparó": FECHA, "En juego": ARS}, {"Regla": 32, "Responsable": 22, "Clave": 18})
last = len(rows) + 1
r0 = last + 3
sh.cell(r0 - 1, 1, "Resumen por regla (fórmulas sobre la tabla de arriba)").font = Font(name=FONT, bold=True, size=11)
for j, h in enumerate(["Regla", "Alertas", "En juego"], 1):
    c = sh.cell(r0, j, h); c.font, c.fill = H_FONT, H_FILL
for k, (rid, (nom, _)) in enumerate(REGLAS.items(), 1):
    sh.cell(r0 + k, 1, nom).font = B_FONT
    sh.cell(r0 + k, 2, f'=COUNTIF($A$2:$A${last},A{r0 + k})').font = B_FONT
    c = sh.cell(r0 + k, 3, f'=SUMIF($A$2:$A${last},A{r0 + k},$G$2:$G${last})'); c.font, c.number_format = B_FONT, ARS
t = r0 + len(REGLAS) + 1
sh.cell(t, 1, "Total").font = BOLD
sh.cell(t, 2, f"=SUM(B{r0 + 1}:B{t - 1})").font = BOLD
c = sh.cell(t, 3, f"=SUM(C{r0 + 1}:C{t - 1})"); c.font, c.number_format = BOLD, ARS

# ---------------------------------------------------------------- A1 · comprobación con fórmulas
a1 = A[A.regla == "equipo_sin_abono"].merge(C["equipo_sin_abono"], left_on="clave", right_on="id_equipo")
rows = [[r.id_equipo, r.id_cliente_x, r.tipo_equipo, r.instalacion, r.abono, r.meses_sin, None, None] for r in a1.itertuples()]
sh = hoja("A1 Sin abono", ["Equipo", "Cliente", "Tipo", "Instalación", "Abono mensual", "Meses sin facturar", "En juego", "Abonos facturados en octubre"],
          rows, {"Instalación": FECHA, "Abono mensual": ARS, "En juego": ARS}, {"Tipo": 22}, {"Abono mensual": IN_FONT, "Meses sin facturar": IN_FONT})
for i in range(2, len(rows) + 2):
    sh[f"G{i}"] = f"=E{i}*F{i}"; sh[f"G{i}"].number_format = ARS
    sh[f"H{i}"] = f'=COUNTIFS(Abonos!$B$2:$B${nA},A{i},Abonos!$A$2:$A${nA},">="&DATE(2025,10,1))'; sh[f"H{i}"].font = LINK_FONT
sh["H1"].comment = Comment("Tiene que dar 0: es la condición de la regla (falta el abono del último mes).", "Matías Garnica")

# ---------------------------------------------------------------- A5 · comprobación con fórmulas
w = C["cobro_sin_visita"]
w = w[w.id_reparto.isin(A[A.regla == "cobro_sin_visita"].clave) & (w.fecha > CORTE - PARAMS["csv_ventana"] * DIA)]
rows = [[r.fecha, r.id_cliente, r.id_reparto, r.importe, None, r.dias_desde_visita] for r in w.itertuples()]
sh = hoja("A5 Cobros sin visita", ["Fecha del cobro", "Cliente", "Reparto", "Importe", "Visitas efectivas ese día", "Días desde la última visita"],
          rows, {"Fecha del cobro": FECHA, "Importe": ARS})
for i in range(2, len(rows) + 2):
    sh[f"E{i}"] = f'=COUNTIFS(Visitas!$B$2:$B${nVi},B{i},Visitas!$A$2:$A${nVi},A{i},Visitas!$D$2:$D${nVi},"Sí")'; sh[f"E{i}"].font = LINK_FONT
sh["E1"].comment = Comment("Tiene que dar 0: el reparto cargó el cobro en efectivo un día en que no visitó al cliente.", "Matías Garnica")

# ---------------------------------------------------------------- A7 · saldo al cambio de ruta, por fórmula
a7 = C["cambio_ruta"]
a7 = a7[a7.id_cliente.astype(str).isin(A[A.regla == "cambio_ruta"].clave)]
rows = [[r.fecha, r.id_cliente, r.reparto_origen, r.reparto_destino, r.motivo, None, r.saldo] for r in a7.itertuples()]
sh = hoja("A7 Cambios de ruta", ["Fecha", "Cliente", "Reparto de origen", "Reparto de destino", "Motivo", "Saldo al cambio (fórmula)", "Saldo según el motor"],
          rows, {"Fecha": FECHA, "Saldo al cambio (fórmula)": ARS, "Saldo según el motor": ARS}, {"Motivo": 28}, {"Saldo según el motor": IN_FONT})
for i in range(2, len(rows) + 2):
    S = lambda hoja_, n, col, extra="": f'SUMIFS({hoja_}!${col}$2:${col}${n},{hoja_}!$B$2:$B${n},$B{i},{hoja_}!$A$2:$A${n},"<="&$A{i}{extra})'
    tipo = lambda t: f',Movimientos!$D$2:$D${nM},"{t}"'
    sh[f"F{i}"] = (f"={S('Ventas', nV, 'D')}+{S('Movimientos', nM, 'E', tipo('Saldo inicial'))}+{S('Movimientos', nM, 'E', tipo('Reversa de pago'))}"
                   f"-{S('Pagos', nP, 'D')}-{S('Movimientos', nM, 'E', tipo('Nota de crédito'))}-{S('Movimientos', nM, 'E', tipo('Incobrable'))}")
    sh[f"F{i}"].font, sh[f"F{i}"].number_format = LINK_FONT, ARS
sh["F1"].comment = Comment("Ventas + saldo inicial + reversas − pagos − notas de crédito − incobrables, hasta el día del cambio. Tiene que coincidir con la columna G.", "Matías Garnica")

# ---------------------------------------------------------------- las otras reglas, con sus datos
a2 = A[A.regla == "equipo_sin_consumo"].merge(C["equipo_sin_consumo"], left_on="clave", right_on="id_equipo")
df_hoja("A2 Sin consumo", pd.DataFrame({"Equipo": a2.id_equipo, "Cliente": a2.id_cliente_x, "Tipo": a2.tipo_equipo, "Instalación": a2.instalacion,
                                         "Última compra con cargo": a2.ult_compra, "Días sin compras": a2.dias_sin, "Valor de reposición": a2.valor,
                                         "Pasado a incobrable": ["Sí" if c in set(b["fact_movimiento"].query("tipo == 'Incobrable'").id_cliente) else "No" for c in a2.id_cliente_x]}),
        {"Instalación": FECHA, "Última compra con cargo": FECHA, "Valor de reposición": ARS}, {"Tipo": 22})
sc = C["solo_sin_cargo"]
a3 = sc[sc.id_cliente.isin(A[A.regla == "solo_sin_cargo"].id_cliente) & (sc.fecha > CORTE - (PARAMS["sc_ventana"] + PARAMS["sc_hist"]) * DIA)]
df_hoja("A3 Solo sin cargo", pd.DataFrame({"Cliente": a3.id_cliente, "Fecha": a3.fecha, "Importe con cargo": a3.importe, "Unidades sin cargo": a3.unidades_sin_cargo,
                                            "En la ventana": np.where(a3.fecha > CORTE - PARAMS["sc_ventana"] * DIA, "Sí", "No")}),
        {"Fecha": FECHA, "Importe con cargo": ARS})
inc = C["incobrables"]
a4 = inc[inc.id_reparto.isin(A[A.regla == "incobrables"].clave) & (inc.fecha > CORTE - PARAMS["inc_ventana"] * DIA)]
df_hoja("A4 Incobrables", pd.DataFrame({"Fecha": a4.fecha, "Cliente": a4.id_cliente, "Reparto": a4.id_reparto, "Importe": a4.importe, "Motivo": a4.motivo,
                                         "Siguió comprando": np.where(a4.sigue_comprando == 1, "Sí", "No")}), {"Fecha": FECHA, "Importe": ARS}, {"Motivo": 24})
nc = C["nc_fuera_patron"].assign(clave=lambda x: x.id_cliente.astype(str) + "|" + x.fecha.dt.strftime("%Y-%m-%d"))
a6 = nc[nc.clave.isin(A[A.regla == "nc_fuera_patron"].clave)]
sh = df_hoja("A6 NC fuera de patrón", pd.DataFrame({"Fecha": a6.fecha, "Cliente": a6.id_cliente, "Reparto": a6.id_reparto, "Importe": a6.importe,
                                                     "Facturación mensual habitual": a6.fac6, "Veces": None, "Motivo cargado": a6.motivo, "Revisión": a6.revision.replace("", "Sin revisar")}),
             {"Fecha": FECHA, "Importe": ARS, "Facturación mensual habitual": ARS}, {"Motivo cargado": 24, "Revisión": 40})
for i in range(2, len(a6) + 2):
    sh[f"F{i}"] = f'=IF(E{i}>0,D{i}/E{i},"")'; sh[f"F{i}"].number_format = '0.0"×"'

# ---------------------------------------------------------------- datos crudos al final
df_hoja("Ventas", ven, {"Fecha": FECHA, "Importe": ARS})
df_hoja("Pagos", pag, {"Fecha": FECHA, "Importe": ARS})
df_hoja("Movimientos", mov, {"Fecha": FECHA, "Importe": ARS}, {"Motivo": 32, "Tipo": 18})
df_hoja("Visitas", vis, {"Fecha": FECHA})
df_hoja("Abonos", abo, {"Fecha": FECHA, "Importe": ARS})
eqd = b["fact_equipo"].merge(b["dim_equipo"], on="id_equipo")[["fecha", "tipo", "id_equipo", "id_cliente", "tipo_equipo", "modalidad", "valor_reposicion", "abono_mensual"]]
eqd.columns = ["Fecha", "Movimiento", "Equipo", "Cliente", "Tipo de equipo", "Modalidad", "Valor de reposición", "Abono mensual"]
df_hoja("Equipos", eqd, {"Fecha": FECHA, "Valor de reposición": ARS, "Abono mensual": ARS}, {"Tipo de equipo": 22, "Modalidad": 22})
cr = b["fact_cambio_reparto"].copy(); cr.columns = ["Fecha", "Cliente", "Reparto de origen", "Reparto de destino", "Motivo"]
df_hoja("Cambios de ruta", cr, {"Fecha": FECHA}, {"Motivo": 30})
cl = b["dim_cliente"][["id_cliente", "nombre", "tipo_cliente", "id_reparto", "id_sucursal", "condicion", "medio_habitual", "frecuencia", "cuenta_interna"]]
cl.columns = ["Cliente", "Nombre", "Tipo", "Reparto actual", "Sucursal", "Condición", "Medio habitual", "Frecuencia", "Cuenta interna"]
df_hoja("Clientes", cl)

# ---------------------------------------------------------------- léeme
txt = [("Base de alertas · Reparto Federal S.A.", "t"),
       (f"Alertas del {CORTE:%d/%m/%Y} con los parámetros acordados. Datos 100% ficticios (generate_alertas.py, semilla fija).", ""), ("", ""),
       ("Hojas", "h"),
       ("Alertas: una fila por alerta (incluye las que el tablero marca como justificadas) y un resumen por regla con fórmulas COUNTIF y SUMIF.", ""),
       ("A1, A5 y A7: la condición de la regla comprobada con fórmulas sobre los datos crudos (abonos de octubre = 0, visitas efectivas el día del cobro = 0, saldo al cambio de ruta).", ""),
       ("A2, A3, A4 y A6: los datos de cada alerta (en A6, las veces sobre la facturación habitual por fórmula).", ""),
       ("Ventas, Pagos, Movimientos, Visitas, Abonos, Equipos, Cambios de ruta y Clientes: los datos crudos.", ""), ("", ""),
       ("Colores", "h"), ("Azul: valor cargado. Verde: fórmula que lee otra hoja. Negro: fórmula de la misma hoja.", ""), ("", ""),
       ("Reglas y parámetros", "h")] + [(f"{REGLAS[k][0]} · responsable: {REGLAS[k][1]}", "") for k in REGLAS] + \
      [("", ""), ("Los criterios completos están en reglas-de-negocio.md y las consultas de SQL Server en la carpeta sql.", "")]
for t_, s_ in txt:
    ws0.append([t_])
    ws0.cell(ws0.max_row, 1).font = Font(name=FONT, bold=True, size=15, color="13294B") if s_ == "t" else BOLD if s_ == "h" else B_FONT
ws0.column_dimensions["A"].width = 150
ws0.sheet_view.showGridLines = False
wb.properties.creator = "Matías Garnica"
wb.properties.title = "Base de alertas · Reparto Federal S.A. (datos ficticios)"
path = os.path.join(DAT, "Base-alertas.xlsx")
wb.save(path)
print("ok ->", os.path.relpath(path), os.path.getsize(path) // 1024, "KB ·", wb.sheetnames)
